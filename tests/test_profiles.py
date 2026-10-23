"""Profiles and scoring rules of server commit f7b795a, on synthetic data only."""
from __future__ import annotations

import io
import json
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

import httpx

from icea_client.api_client import ApiClient
from icea_client.compare import run_compare, select_record
from icea_client.config import ApiSettings
from icea_client.ground_truth_uat import load_ground_truth
from icea_client.preflight import run_preflight
from icea_client.runner import RunRefused, run
from icea_client.workbook import error_analysis
from tests.fixtures import FakeApi, environment, field, response_body
from contracts import CANONICAL_PROFILES, CONTRACTS, LEGACY_PROFILES
from normalization import date
from schemas import validate_payload_with_field_issues

SEL = "subcontractor_selection.entries[]."
TOKEN = "synthetic-token"


def quiet(*_args, **_kwargs):
    pass


# (row_id, field_path, role, page, group, raw, normalized, status, rule)
SELECTION_ROWS = [
    ("S1", "subcontractor_selection.document_date", "COND", "1", None, "05/10/23", "2023-10-05", "Presente", "date_equivalent"),
    ("S2", SEL + "operation", "CORE", "1", "1", "Tintura", "Tintura", "Presente", "grouped_value"),
    # One quantity without direction: the same value on both fields.
    ("S3", SEL + "quantity_input_kg", "COND", "1", "1", "120", 120, "Presente", "numeric_equivalent"),
    ("S4", SEL + "quantity_output_kg", "COND", "1", "1", "120", 120, "Presente", "numeric_equivalent"),
    ("S5", SEL + "operation", "CORE", "1", "2", "Filatura", "Filatura", "Presente", "grouped_value"),
    # Explicit direction: input only.
    ("S6", SEL + "quantity_input_kg", "COND", "1", "2", "50", 50, "Presente", "numeric_equivalent"),
    ("S7", SEL + "quantity_output_kg", "COND", "1", "2", None, None, "Assente", "numeric_equivalent"),
    ("S8", SEL + "subcontractor", "CORE", "1", "2", "Terzista Sintetico", "Terzista Sintetico", "Ambiguo", "grouped_value"),
]
SELECTION_FIELDS = [
    field("subcontractor_selection.document_date", "05/10/23"),
    field(SEL + "operation", "Tintura", group="A"),
    field(SEL + "quantity_input_kg", "120", group="A"),
    field(SEL + "quantity_output_kg", "120", group="A"),
    field(SEL + "operation", "Filatura", group="B"),
    field(SEL + "quantity_input_kg", "50", group="B"),
    # Copied into the other direction although the document states the input only.
    field(SEL + "quantity_output_kg", "50", group="B"),
    field(SEL + "subcontractor", "Terzista Sintetico", group="B"),
]
RMDF_ROWS = [
    ("F1", "invoice.seller.name", "CORE", "1", None, "Fornitore Sintetico", "Fornitore Sintetico", "Presente", "text_normalized"),
    ("F2", "invoice.seller.address.country", "COND", "1", None, "IT", "Italia", "Presente", "text_normalized"),
    ("F3", "invoice.number", "CORE", "1", None, "R-77", "R-77", "Presente", "identifier_exact"),
    ("F4", "invoice.date", "CORE", "1", None, "2024-03-01", "2024-03-01", "Presente", "date_equivalent"),
    ("F5", "invoice.lines[].lot_numbers[]", "CORE", "1", "1", "LT-1", "LT-1", "Presente", "identifier_exact"),
    ("F6", "invoice.lines[].quantity", "CORE", "1", "1", "1234,56", 1234.5600000000002, "Presente", "numeric_equivalent"),
    ("F7", "invoice.customer.name", "CORE", None, None, "valore", None, "Assente", "text_normalized"),
]
RMDF_FIELDS = [
    field("invoice.seller.name", "Fornitore Sintetico"),
    field("invoice.seller.address.country", "IT"),
    field("invoice.number", "R-77"),
    field("invoice.date", "01/03/2024"),
    field("invoice.lines[].lot_numbers[]", "LT-1", group="L1"),
    field("invoice.lines[].quantity", "1234.56", group="L1"),
]
SALE_ROWS = [
    ("V1", "invoice.number", "CORE", "1", None, "V-1", "V-1", "Presente", "identifier_exact"),
    ("V2", "invoice.lines[].lot_numbers[]", "CORE", "1", "1", "LOT-X", "LOT-X", "Presente", "identifier_exact"),
    ("V3", "invoice.lines[].net_weight", "CORE", "1", "1", "100", 100, "Presente", "numeric_equivalent"),
    ("V4", "invoice.lines[].gross_weight", "COND", None, "1", None, None, "Assente", "numeric_equivalent"),
]
SALE_FIELDS = [
    field("invoice.number", "V-1"),
    field("invoice.lines[].lot_numbers[]", "LOT-X", group="1"),
    # Correct quantity reported in a sibling field of the same line.
    field("invoice.lines[].gross_weight", "100", group="1"),
]


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def campaign(self, samples, profiles, fields_by_sample, responder=None):
        base, workbook, hashes = environment(self.root, samples, profile=profiles)
        gt = load_ground_truth(workbook)

        def respond(request, form):
            sample_id = form["filename"].removeprefix("synthetic_").removesuffix(".pdf")
            if responder:
                result = responder(request, form, hashes[sample_id])
                if result is not None:
                    return result
            return httpx.Response(200, json=response_body(form, filename=form["filename"],
                                                          sha256=hashes[sample_id],
                                                          fields=fields_by_sample[sample_id]))

        api = FakeApi(hashes, respond)
        settings = ApiSettings("http://localhost:8000", TOKEN, 5)
        selected = list(samples)
        run(gt, base, selected, execute=True, approved_calls=4 * len(selected),
            client_factory=lambda: ApiClient(settings, transport=httpx.MockTransport(api)),
            base_url="http://localhost:8000", log=quiet)
        with redirect_stdout(io.StringIO()):
            result = run_compare(gt, base, log=quiet)
        return result, api, base, gt


class ApiProfileTests(Base):
    def test_profile_set_matches_server(self):
        self.assertEqual(CANONICAL_PROFILES, ("TC", "FATTURA-VENDITA", "DDT-VENDITA", "FATTURA-ACQUISTO-RMDF",
                                              "DDT-ACQUISTO-RMDF", "SELEZIONE-TERZISTI"))
        self.assertEqual(len(CONTRACTS["FATTURA-ACQUISTO-RMDF"]), 14)
        self.assertEqual(len(CONTRACTS["DDT-ACQUISTO-RMDF"]), 15)
        self.assertNotIn("transport_document.handwritten_order_or_lot", CONTRACTS["DDT-ACQUISTO-RMDF"])
        selection = CONTRACTS["SELEZIONE-TERZISTI"]
        self.assertEqual(len(selection), 15)
        self.assertTrue(all(path.startswith("subcontractor_selection.") for path in selection))
        self.assertNotIn(SEL + "quantity_kg", selection)

    def test_legacy_profiles_are_blocked_before_the_api_and_rejected_by_it(self):
        rows = [("E1", "subcontractor_process.production_lot", "CORE", "1", None, "P1", "P1",
                 "Presente", "identifier_exact")]
        base, workbook, hashes = environment(self.root, {"UAT-01": rows, "UAT-02": SALE_ROWS},
                                             profile={"UAT-01": "ELENCO-TERZISTI",
                                                      "UAT-02": "FATTURA-VENDITA"})
        gt = load_ground_truth(workbook)
        report = run_preflight(gt, base)
        self.assertEqual(report["excluded_samples"][0]["reason"], "legacy_profile_rejected_by_api")
        explicit = run_preflight(gt, base, ["UAT-01"])
        self.assertEqual(explicit["samples"][0]["blocking"], ["legacy_profile_rejected_by_api"])
        with self.assertRaises(RunRefused):
            run(gt, base, ["UAT-01"], log=quiet)
        api = FakeApi(hashes)
        settings = ApiSettings("http://localhost:8000", TOKEN, 5)
        with ApiClient(settings, transport=httpx.MockTransport(api)) as client:
            for profile in LEGACY_PROFILES:
                record = client.extract(base.samples / "synthetic_UAT-01.pdf", "synthetic_UAT-01.pdf",
                                        profile, "A")
                self.assertEqual((record["http_status"], record["error"]), (422, "http_422"))


class SelectionTests(Base):
    def test_single_row_header_preflight_and_ambiguous_rows(self):
        rows = [*SELECTION_ROWS,
                ("S9", SEL + "quantity_kg", "COND", "1", "1", "120", 120, "Presente", "numeric_equivalent"),
                ("S10", SEL + "lot_number", "COND", "1", "1", "L-9", "L-9", "Assente", "identifier_exact")]
        base, workbook, hashes = environment(self.root, {"UAT-07": rows}, profile="SELEZIONE-TERZISTI")
        document = load_ground_truth(workbook).documents["UAT-07"]
        self.assertEqual((document.filename, document.sha256), ("synthetic_UAT-07.pdf", hashes["UAT-07"]))
        sample = run_preflight(load_ground_truth(workbook), base)["samples"][0]
        self.assertEqual(sample["blocking"], [])
        self.assertEqual(sample["issues"]["unknown_field_path"]["row_ids"], ["S9"])
        self.assertEqual(sample["issues"]["absent_with_value"]["row_ids"], ["S10"])
        self.assertEqual(sample["explicit_ambiguities"]["row_ids"], ["S8"])

    def test_duplicated_and_directional_quantities(self):
        result, _, _, _ = self.campaign({"UAT-07": SELECTION_ROWS}, "SELEZIONE-TERZISTI",
                                        {"UAT-07": SELECTION_FIELDS})
        report = result["data"]["samples"][0]["methods"]["A"]["score"]
        outcomes = {row["row_id"]: row["outcome"] for row in report["rows"]}
        self.assertEqual([outcomes[key] for key in ("S1", "S2", "S3", "S4", "S5", "S6")], ["matched"] * 6)
        self.assertEqual(outcomes["S7"], "incorrect")  # direction stated: no copy allowed
        self.assertNotIn("S8", outcomes)  # Ambiguo: excluded from the score
        self.assertEqual(report["ambiguous_truth_rows"], 1)
        reasons = {item["row_id"]: item["reasons"] for item in report["not_ready"]}
        self.assertEqual(reasons["S8"], ["ambiguous_in_ground_truth"])

    def test_per_group_cardinality_becomes_field_issue(self):
        payload = {"processing_status": "partial", "warnings": [], "fields": [
            field(SEL + "operation", "Tintura", group="A"),
            field(SEL + "quantity_input_kg", "10", group="A"),
            field(SEL + "quantity_input_kg", "20", group="A")]}
        extraction, issues = validate_payload_with_field_issues(
            payload, source_file="x.pdf", source_sha256="0" * 64, profile="SELEZIONE-TERZISTI",
            strategy="vision_direct", page_count=1)
        self.assertEqual(len(extraction.fields), 1)
        self.assertEqual((issues[0]["label"], issues[0]["field_path"], issues[0]["group_id"]),
                         ("collisione_cardinalita", SEL + "quantity_input_kg", "A"))
        rmdf = {"processing_status": "partial", "warnings": [], "fields": [
            field("invoice.lines[].quantity", "1", group="L1"),
            field("invoice.lines[].quantity", "2", group="L1")]}
        _, issues = validate_payload_with_field_issues(
            rmdf, source_file="x.pdf", source_sha256="0" * 64, profile="FATTURA-ACQUISTO-RMDF",
            strategy="vision_direct", page_count=1)
        self.assertEqual(issues[0]["field_path"], "invoice.lines[].quantity")

    def collision_campaign(self, rows, profile, fields, issue, root):
        def responder(request, form, sha256):
            if form["extraction_option"] == "A":
                return httpx.Response(200, json=response_body(
                    form, filename=form["filename"], sha256=sha256, fields=fields,
                    field_issues=[issue]))
            return None

        self.root = root
        result, _, _, _ = self.campaign({"UAT-07": rows}, profile, {"UAT-07": fields}, responder)
        report = result["data"]["samples"][0]["methods"]["A"]["score"]
        return report, {row["row_id"]: row["outcome"] for row in report["rows"]}, result

    def test_collision_is_partial_when_one_value_matches_else_incorrect(self):
        fields = [item for item in SELECTION_FIELDS if not (item["field_path"].endswith("input_kg")
                                                           and item["group_id"] == "A")]
        base_root = self.root
        for display, expected in (("'120'; '200'", "partial"), ("'7'; '8'", "incorrect")):
            issue = {"label": "collisione_cardinalita", "field_path": SEL + "quantity_input_kg",
                     "group_id": "A", "page": 1, "value_display": display, "source_field_indices": [2, 3]}
            report, outcomes, result = self.collision_campaign(
                SELECTION_ROWS, "SELEZIONE-TERZISTI", fields, issue, base_root / expected)
            self.assertEqual((report["invalid_field"], outcomes["S3"]), (0, expected))
            row = next(row for row in result["data"]["samples"][0]["rows"] if row["row_id"] == "S3")
            self.assertEqual(row["methods"]["A"]["outcome"], expected)


GT_LIST_ROWS = [
    ("V1", "invoice.number", "CORE", "1", None, "V-1", "V-1", "Presente", "identifier_exact"),
    ("V2", "invoice.lines[].lot_numbers[]", "CORE", "1", "1", "LOT-X", "LOT-X", "Presente", "identifier_exact"),
    # Cell typed as '12'; '15' in Excel: the leading apostrophe is hidden, the text reads 12'; '15'.
    ("V3", "invoice.lines[].measurement_value", "COND", "1", "1", "12'; '15'", None, "Presente", "numeric_equivalent"),
]


class GroundTruthListTests(Base):
    def run_case(self, name, fields, issues=None):
        def responder(request, form, sha256):
            return httpx.Response(200, json=response_body(form, filename=form["filename"], sha256=sha256,
                                                          fields=fields, field_issues=issues or []))

        self.root = self.root / name
        result, _, _, gt = self.campaign({"UAT-03": GT_LIST_ROWS}, "FATTURA-VENDITA",
                                         {"UAT-03": fields}, responder)
        report = result["data"]["samples"][0]["methods"]["A"]["score"]
        return gt, {row["row_id"]: row["outcome"] for row in report["rows"]}

    def test_hidden_apostrophe_list_is_read_and_scored(self):
        base = [field("invoice.number", "V-1"), field("invoice.lines[].lot_numbers[]", "LOT-X", group="1")]
        both = {"label": "collisione_cardinalita", "field_path": "invoice.lines[].measurement_value",
                "group_id": "1", "page": 1, "value_display": "'15'; '12'", "source_field_indices": [2, 3]}
        gt, outcomes = self.run_case("all", base, [both])
        self.assertEqual(gt.documents["UAT-03"].rows[2].value_raw, "12'; '15'")
        quoted = __import__("evaluation").quoted_values
        self.assertEqual(quoted("12'; '15'"), ["12", "15"])
        self.assertEqual(outcomes["V3"], "matched")
        _, outcomes = self.run_case("one", [*base, field("invoice.lines[].measurement_value", "12", group="1")])
        self.assertEqual(outcomes["V3"], "partial")
        _, outcomes = self.run_case("none", [*base, field("invoice.lines[].measurement_value", "99", group="1")])
        self.assertEqual(outcomes["V3"], "incorrect")

class ScoringRuleTests(Base):
    def test_rmdf_country_variants_excel_precision_and_absent_with_value(self):
        result, _, _, gt = self.campaign({"UAT-05": RMDF_ROWS}, "FATTURA-ACQUISTO-RMDF",
                                         {"UAT-05": RMDF_FIELDS})
        quantity = next(row for row in gt.documents["UAT-05"].rows if row.row_id == "F6")
        self.assertEqual(quantity.value_normalized, "1234.56")
        report = result["data"]["samples"][0]["methods"]["B"]["score"]
        outcomes = {row["row_id"]: row["outcome"] for row in report["rows"]}
        self.assertEqual(set(outcomes.values()), {"matched"})
        self.assertNotIn("F7", outcomes)
        self.assertEqual({item["row_id"]: item["reasons"] for item in report["not_ready"]},
                         {"F7": ["absent_with_value"]})

    def test_two_digit_year(self):
        self.assertEqual(date("05/10/23"), "2023-10-05")

    def test_alternative_quantity_is_a_match_and_not_an_error(self):
        result, _, _, _ = self.campaign({"UAT-03": SALE_ROWS}, "FATTURA-VENDITA", {"UAT-03": SALE_FIELDS})
        sample = result["data"]["samples"][0]
        report = sample["methods"]["C"]["score"]
        outcomes = {row["row_id"]: row for row in report["rows"]}
        self.assertEqual(outcomes["V3"]["outcome"], "matched")
        self.assertEqual(outcomes["V3"]["alternative_field_path"], "invoice.lines[].gross_weight")
        self.assertEqual(outcomes["V4"]["outcome"], "matched")
        stats, errors = error_analysis(result["data"])
        self.assertEqual(errors, [])
        self.assertTrue(all(sum(stat["errors"]) == 0 for stat in stats))

    def test_schema_failure_is_a_run_failure_excluded_from_error_metrics(self):
        def responder(request, form, sha256):
            if form["extraction_option"] == "D":
                return httpx.Response(502, json={"error": "schema", "run_id": "r"})
            return None

        result, _, _, _ = self.campaign({"UAT-03": SALE_ROWS}, "FATTURA-VENDITA",
                                        {"UAT-03": [field("invoice.number", "V-9")]}, responder)
        sample = result["data"]["samples"][0]
        self.assertIsNone(sample["methods"]["D"]["score"])
        stats, errors = error_analysis(result["data"])
        option_d = 3
        self.assertTrue(all(stat["errors"][option_d] == 0 for stat in stats))
        self.assertTrue(errors and all(row[11 + 4 * option_d] == "errore API: schema" for row in errors))
        self.assertEqual(result["summary"]["totals"]["D"]["valid_samples"], 0)

    def test_latest_schema_failure_wins_over_older_success_but_not_transient_errors(self):
        class Document:
            profile, sha256 = "TC", "a" * 64

        def record(status, started, error=None, http=200):
            return {"method": "A", "extraction_option": "A", "api_version": "2.0",
                    "profile": "TC", "pdf_sha256": "a" * 64, "status": status,
                    "started_at": started, "error": error, "http_status": http}

        ok = record("succeeded", "2026-10-01")
        schema = record("failed", "2026-10-02", "schema", 502)
        busy = record("failed", "2026-10-03", "busy", 429)
        self.assertIs(select_record([ok, schema, busy], Document, "A"), schema)
        self.assertIs(select_record([ok, busy], Document, "A"), ok)
        self.assertIs(select_record([busy], Document, "A"), busy)

    def test_reports_keep_code_and_gt_hashes_without_server_details(self):
        result, _, base, _ = self.campaign({"UAT-03": SALE_ROWS}, "FATTURA-VENDITA", {"UAT-03": SALE_FIELDS})
        summary = json.loads((base.reports / "comparison_summary.json").read_text(encoding="utf-8"))
        self.assertEqual(len(summary["code"]["canonical_files_sha256"]), 7)
        self.assertEqual(len(summary["code"]["client_sources_sha256"]), 64)
        self.assertEqual(len(summary["ground_truth_sha256"]), 64)
        method = summary["samples"][0]["methods"]["A"]
        self.assertFalse({"prompt_sha256", "model_version"} & set(method))
        self.assertEqual(set(method["timing"]), {"pipeline_seconds", "http_seconds"})
        from openpyxl import load_workbook
        names = load_workbook(result["workbook"]).sheetnames
        self.assertEqual(names[:3], ["Sintesi", "Errori per campo", "Errori"])


TC_BOOLEAN_ROWS = [
    ("T1", "tc.number", "CORE", "1", None, "TC-1", "TC-1", "Presente", "identifier_exact"),
    ("T2", "tc.is_draft", "CORE", None, None, "false", "false", "Presente", "boolean_judgment"),
    ("T3", "tc.usda_nop", "COND", None, None, "true", "true", "Presente", "boolean_judgment"),
]


class ImplicitFalseTests(Base):
    def test_missing_boolean_counts_as_false_but_ambiguous_does_not(self):
        fields = {"A": [field("tc.number", "TC-1")],
                  "B": [field("tc.number", "TC-1"),
                        field("tc.usda_nop", None, page=None, status="ambiguous")]}

        def responder(request, form, sha256):
            key = "B" if form["extraction_option"] in {"B", "D"} else "A"
            return httpx.Response(200, json=response_body(form, filename=form["filename"],
                                                          sha256=sha256, fields=fields[key]))

        result, _, _, _ = self.campaign({"UAT-01": TC_BOOLEAN_ROWS}, "TC", {"UAT-01": []}, responder)
        sample = result["data"]["samples"][0]
        option_a = {row["row_id"]: row["outcome"] for row in sample["methods"]["A"]["score"]["rows"]}
        option_b = {row["row_id"]: row["outcome"] for row in sample["methods"]["B"]["score"]["rows"]}
        self.assertEqual((option_a["T2"], option_a["T3"]), ("matched", "incorrect"))
        self.assertEqual(sample["methods"]["A"]["implicit_false_paths"], ["tc.is_draft", "tc.usda_nop"])
        self.assertEqual(sample["methods"]["B"]["implicit_false_paths"], ["tc.is_draft"])
        self.assertEqual((option_b["T2"], option_b["T3"]), ("matched", "omitted"))
        from openpyxl import load_workbook
        sheet = load_workbook(result["workbook"])["UAT-01"]
        values = {sheet.cell(row=row, column=3).value: sheet.cell(row=row, column=11).value
                  for row in range(16, 19)}
        self.assertEqual((values["T2"], values["T3"]), ("false", "false"))

RMDF_DDT_ROWS = [
    ("D1", "transport_document.sender.name", "CORE", "1", None, "Mittente Sintetico", "Mittente Sintetico", "Presente", "text_normalized"),
    ("D2", "transport_document.number", "CORE", "1", None, "D-5", "D-5", "Presente", "identifier_exact"),
    ("D3", "transport_document.lines[].lot_numbers[]", "CORE", "1", "1", "LT-9", "LT-9", "Presente", "identifier_exact"),
]
RMDF_DDT_FIELDS = [
    field("transport_document.sender.name", "Mittente Sintetico"),
    field("transport_document.number", "D-5"),
    field("transport_document.lines[].lot_numbers[]", "LT-9", group="1"),
]
REMOVED = "transport_document.handwritten_order_or_lot"


class RemovedFieldTests(Base):
    def test_removed_path_is_rejected_by_the_schema(self):
        payload = {"processing_status": "partial", "warnings": [],
                   "fields": [*RMDF_DDT_FIELDS, field(REMOVED, "123", group="h1")]}
        with self.assertRaisesRegex(ValueError, "unknown field_path"):
            validate_payload_with_field_issues(payload, source_file="x.pdf", source_sha256="0" * 64,
                                               profile="DDT-ACQUISTO-RMDF", strategy="vision_direct",
                                               page_count=1)

    def test_archived_run_with_removed_path_is_rescored_without_rewriting(self):
        _, _, base, gt = self.campaign({"UAT-06": RMDF_DDT_ROWS}, "DDT-ACQUISTO-RMDF", {"UAT-06": RMDF_DDT_FIELDS})
        # Simulate a response archived before the field was removed from the contract.
        archived = next(base.runs.glob("UAT-06/A/*/response.json"))
        response = json.loads(archived.read_text(encoding="utf-8"))
        response["body"]["extraction"]["fields"].append(field(REMOVED, "123", group="h1"))
        archived.write_text(json.dumps(response), encoding="utf-8")
        before = archived.read_bytes()
        with redirect_stdout(io.StringIO()):
            result = run_compare(gt, base, log=quiet)
        method = result["data"]["samples"][0]["methods"]["A"]
        self.assertEqual((method["status"], method["removed_paths"]), ("succeeded", [REMOVED]))
        self.assertEqual(method["score"]["matched"], 3)
        self.assertFalse(any(row["field_path"] == REMOVED for row in result["data"]["samples"][0]["rows"]))
        self.assertEqual(archived.read_bytes(), before)

ON_BEHALF = "tc.buyer.buying_on_behalf_of"
FALSE_TEXT_ROWS = [
    ("B1", "tc.number", "CORE", "1", None, "TC-1", "TC-1", "Presente", "identifier_exact"),
    ("B2", ON_BEHALF, "CORE", "1", "1", "false", "false", "Presente", "text_normalized"),
    # A buyer name in the same group lets the scorer align the model buyer group.
    ("B3", "tc.buyer.name", "CORE", "1", "1", "Acquirente Sintetico", "Acquirente Sintetico", "Presente", "text_normalized"),
]


class FalseAsEmptyTests(Base):
    def test_false_text_value_means_empty_field(self):
        number = [field("tc.number", "TC-1"), field("tc.buyer.name", "Acquirente Sintetico", group="1")]
        by_method = {"A": [*number],
                     "B": [*number, field(ON_BEHALF, None, page=None, status="not_present")],
                     "C": [*number, field(ON_BEHALF, "false", group="1")],
                     "D": [*number, field(ON_BEHALF, "Soggetto Sintetico", group="1")]}

        def responder(request, form, sha256):
            return httpx.Response(200, json=response_body(form, filename=form["filename"], sha256=sha256,
                                                          fields=by_method[form["extraction_option"]]))

        result, _, _, gt = self.campaign({"UAT-01": FALSE_TEXT_ROWS}, "TC", {"UAT-01": []}, responder)
        sample = result["data"]["samples"][0]
        outcome = {method: {row["row_id"]: row["outcome"] for row in sample["methods"][method]["score"]["rows"]}["B2"]
                   for method in ("A", "B", "C", "D")}
        self.assertEqual(outcome, {"A": "matched", "B": "matched",
                                   "C": "matched", "D": "incorrect"})
        self.assertEqual(sample["false_as_empty_rows"], ["B2"])
        self.assertEqual(sample["methods"]["C"]["false_as_empty_paths"], [ON_BEHALF])
        # The GT itself is not changed: the loader still reads the original row.
        original = next(row for row in gt.documents["UAT-01"].rows if row.row_id == "B2")
        self.assertEqual((original.status, original.value_raw), ("Presente", "false"))
        shown = next(row for row in sample["rows"] if row.get("row_id") == "B2")
        self.assertEqual(shown["gt_raw"], "false")


if __name__ == "__main__":
    unittest.main()
